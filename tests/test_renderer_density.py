"""Renderer bidirectional density — sparse resumes scale UP to fill the page,
dense resumes keep the shrink tiers. Real fpdf2 measure passes, no LLM.

These assert the SINGLE-PAGE shrink engine (the default/prod contract), so they
force ALLOW_TWO_PAGE off. Two-page behavior is covered by test_two_page.py.
"""
import pytest

from services import resume_renderer
from services.resume_renderer import ResumeRenderer


@pytest.fixture(autouse=True)
def _force_single_page(monkeypatch):
    monkeypatch.setattr(resume_renderer, "ALLOW_TWO_PAGE", False, raising=False)


def _bullet(i):
    return (
        f"Delivered measurable platform improvement number {i} by refactoring the "
        f"ingestion service and tuning PostgreSQL indexes for clinical workloads."
    )


def _resume(n_roles, bullets_per_role, n_projects=0, bullets_per_project=0):
    return {
        "contact": {
            "name": "Surya Kommuguri", "email": "s@example.com",
            "phone": "480-555-0100", "linkedin": "linkedin.com/in/surya",
            "github": "", "location": "Tempe, AZ", "portfolio": "",
        },
        "summary": (
            "Backend engineer with production experience building Python "
            "microservices, REST APIs, and event-driven systems across healthcare "
            "platforms on AWS and Kubernetes."
        ),
        "skills": {
            "Cloud & Infrastructure": ["AWS", "Docker", "Kubernetes", "Helm"],
            "Databases": ["PostgreSQL", "MySQL", "MongoDB"],
            "Languages & Frameworks": ["Python", "FastAPI", "Django", "SQL"],
            "Observability": ["Grafana", "Datadog APM", "ELK Stack", "Kafka"],
        },
        "experience": [
            {
                "title": f"Software Engineer {r}", "company": f"Company {r}",
                "location": "Tempe, AZ",
                "dates": f"January 20{20 + r} - December 20{21 + r}",
                "type": "Full-time",
                "bullets": [_bullet(f"{r}.{b}") for b in range(bullets_per_role)],
            }
            for r in range(n_roles)
        ],
        "education": [{
            "degree": "Master of Science in Information Technology",
            "institution": "Arizona State University", "location": "Tempe, AZ",
            "dates": "January 2024 - December 2025", "gpa": "", "coursework": "",
        }],
        "projects": [
            {
                "name": f"Telemetry Toolkit {p}", "dates": "",
                "bullets": [_bullet(f"p{p}.{b}") for b in range(bullets_per_project)],
                "tech": "Python, Kafka",
            }
            for p in range(n_projects)
        ],
        "certifications": [],
    }


SPARSE = _resume(n_roles=2, bullets_per_role=4)                  # experienced, no projects

# Content-volume STRESS shape (not a realistic persona — the experience-first
# policy means a senior never gets 3 generated projects). Guards the ladder's
# refusal to inflate when the page is already full.
DENSITY_STRESS = _resume(n_roles=3, bullets_per_role=6, n_projects=3, bullets_per_project=4)


def _senior(role_bullet_counts):
    """Realistic senior shape: several roles, recency-weighted bullet depth,
    NO projects section (experience-first policy)."""
    resume = _resume(len(role_bullet_counts), 1)
    for i, n in enumerate(role_bullet_counts):
        resume["experience"][i]["bullets"] = [
            _bullet(f"{i}.{b}") for b in range(n)
        ]
    resume["projects"] = []
    return resume


SENIOR = _senior([6, 5, 4, 3])  # 4 roles, 18 bullets, experience-only


def test_sparse_resume_takes_the_10pt_up_tier():
    """Thin resume takes the largest body that still keeps section_pt 1.5pt
    above the body (10pt). Stretch stays off."""
    r = ResumeRenderer()
    body, lh, lh_s, height = r._choose_up_tier(SPARSE)
    assert body == 10.0
    _, res = r._compute_pdf_layout(SPARSE)
    assert res.tier.body_pt == 10.0
    assert res.stretch == 1.0


def test_density_stress_stays_at_base_or_below():
    r = ResumeRenderer()
    body, lh, lh_s, height = r._choose_up_tier(DENSITY_STRESS)
    assert body <= 10.0, "over-full content must not inflate"


def test_realistic_senior_fills_one_page():
    """4 roles at 6/5/4/3 bullets, no projects — professional density
    (content-rich resumes never inflate) that fills the page without spilling."""
    from services.pdf_ats_check import check_pdf_parseability

    r = ResumeRenderer()
    body, _, _, height = r._choose_up_tier(SENIOR)
    assert body <= 10.0
    assert height / r._AVAIL_H >= 0.80

    pdf, res = r.generate_pdf(SENIOR)
    check = check_pdf_parseability(pdf, SENIOR, {}, outcome=res.outcome)
    assert check["checks"]["page_count"] == 1
    assert check["pdf_parseable_score"] >= 85


def test_up_tier_measures_monotonic():
    """Bigger presentation must measure taller — guards the ladder ordering."""
    r = ResumeRenderer()
    h10 = r._measure_min(SPARSE, 10.0, 4.01, 4.01)
    h94 = r._measure_min(SPARSE, 9.4, 3.77, 3.77)
    assert h10 > h94


@pytest.mark.parametrize(
    "resume,max_pages",
    [(SPARSE, 1), (SENIOR, 1), (DENSITY_STRESS, 2)],
    ids=["sparse", "senior", "stress"],
)
def test_render_stays_parseable(resume, max_pages):
    """Letter is shorter than A4 — stress content may take two pages, never clip."""
    from services.pdf_ats_check import check_pdf_parseability

    pdf_bytes, res = ResumeRenderer().generate_pdf(resume)
    check = check_pdf_parseability(pdf_bytes, resume, {}, outcome=res.outcome)
    assert check["checks"]["page_count"] == res.pages
    assert res.pages <= max_pages
    assert check["checks"]["extraction_ok"] is True
    assert check["checks"]["no_glyph_corruption"] is True
    if res.outcome == "fit_one":
        assert check["pdf_parseable_score"] >= 85


def test_docx_renders_with_spec_font(tmp_path):
    from docx import Document
    import io

    r = ResumeRenderer()
    data = r.generate_docx(SPARSE)
    doc = Document(io.BytesIO(data))
    assert doc.styles["Normal"].font.name == r.spec.docx_font_family
    text = "\n".join(p.text for p in doc.paragraphs)
    # Name renders in Title Case (e.g. "Surya Kommuguri", not all caps).
    assert "Surya Kommuguri" in text
    assert "EXPERIENCE" in text


def test_pdf_embeds_charis():
    """The PDF must carry the embedded Charis SIL family."""
    pdf_bytes, _res = ResumeRenderer().generate_pdf(SPARSE)
    assert b"Charis" in pdf_bytes


def test_pdf_name_renders_title_case():
    """Name renders in Title Case (e.g. "Surya Kommuguri", not all caps) —
    and the parseability checker must still find it in the text layer."""
    from services.pdf_ats_check import extract_pdf_text

    pdf_bytes, _res = ResumeRenderer().generate_pdf(SPARSE)
    text, _ = extract_pdf_text(pdf_bytes)
    assert "Surya Kommuguri" in text


# ---------------------------------------------------------------------------
# Plain-text contacts (product decision: no hyperlinks in either format)
# ---------------------------------------------------------------------------

def test_pdf_contacts_are_plain_text():
    """No link annotations anywhere in the PDF — email/LinkedIn render as
    plain black text."""
    pdf_bytes, _res = ResumeRenderer().generate_pdf(SPARSE)
    assert b"/URI" not in pdf_bytes


def test_docx_contacts_are_plain_text():
    import io
    from docx import Document

    data = ResumeRenderer().generate_docx(SPARSE)
    doc = Document(io.BytesIO(data))
    hyperlink_rels = [
        rel for rel in doc.part.rels.values() if "hyperlink" in rel.reltype
    ]
    assert not hyperlink_rels
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "s@example.com" in text            # email still present, as text
    assert "linkedin.com/in/surya" in text


# ---------------------------------------------------------------------------
# Pass-3 leading stretch — page fill for resumes between the font rungs
# ---------------------------------------------------------------------------

# The dead-band shape this fix exists for: content that just misses the 11pt
# up-tier (like the user's real 2-role 8+6-bullet resume) used to render at
# 10.5pt with ~25% of the page empty because the gap ceilings could only
# absorb ~27mm.
MID_DENSITY = _resume(n_roles=2, bullets_per_role=7)


def test_mid_density_fills_page():
    r = ResumeRenderer()
    _working, layout = r._compute_pdf_layout(MID_DENSITY)
    fill = layout.fill_ratio
    assert 0.90 <= fill <= r.spec.max_fill + 1e-9, f"fill {fill:.1%}"
    assert layout.stretch <= r._STRETCH_CAP + 1e-9
    assert layout.tier.lh * layout.stretch <= r._MAX_LEADING_MM + 1e-6


def test_sparse_resume_renders_compactly_not_padded():
    """Thin resumes get larger type and wider margins, not stretched leading."""
    r = ResumeRenderer()
    _working, layout = r._compute_pdf_layout(SPARSE)
    assert layout.tier.body_pt == 10.0
    assert layout.stretch == 1.0
    assert layout.gaps.margin_lr == pytest.approx(r.spec.margin_lr_wide_mm)
    assert layout.fill_ratio <= r.spec.max_fill + 1e-9


def test_fill_ramp_leaves_fillable_pages_alone():
    """The ramp must only bite on pages that CANNOT be filled honestly — the
    dead-band and senior shapes still reach the bottom."""
    r = ResumeRenderer()
    for resume in (MID_DENSITY, SENIOR):
        _working, layout = r._compute_pdf_layout(resume)
        assert layout.fill_ratio <= r.spec.max_fill + 1e-9
        assert layout.fill_ratio >= 0.88


def test_heading_and_job_gaps_stay_tight():
    """Pass-2 leftover must not open holes under headings, between roles,
    or between sections (SUMMARY → EXPERIENCE, etc.)."""
    r = ResumeRenderer()
    for resume in (SPARSE, MID_DENSITY, SENIOR, DENSITY_STRESS):
        _working, layout = r._compute_pdf_layout(resume)
        assert layout.gaps.post_header >= r._MIN_POST_HEADER - 1e-9
        assert layout.gaps.post_header <= r._MAX_POST_HEADER + 1e-9
        assert layout.gaps.entry_gap <= r._MAX_ENTRY_GAP + 1e-9
        assert layout.gaps.section_gap <= r._MAX_SECTION_GAP + 1e-9
        assert layout.gaps.skill_gap >= r._MIN_SKILL_GAP - 1e-9
        assert layout.gaps.skill_gap <= r._MAX_SKILL_GAP + 1e-9


def test_dense_resume_never_stretches():
    r = ResumeRenderer()
    _working, layout = r._compute_pdf_layout(DENSITY_STRESS)
    assert layout.stretch == pytest.approx(1.0, abs=0.02)
    assert layout.gaps.post_header >= r._MIN_POST_HEADER - 1e-9


def test_stretch_still_single_page():
    """Stretched renders must not spill to page 2."""
    from services.pdf_ats_check import check_pdf_parseability

    pdf_bytes, res = ResumeRenderer().generate_pdf(MID_DENSITY)
    check = check_pdf_parseability(pdf_bytes, MID_DENSITY, {}, outcome=res.outcome)
    assert check["checks"]["page_count"] == 1
    assert check["pdf_parseable_score"] >= 85


# ---------------------------------------------------------------------------
# Skills block spacing
# ---------------------------------------------------------------------------

_MEASURE_KNOBS = dict(
    section_gap=1.2, entry_gap=0.9, post_header=1.0, header_gap=0.5,
    body_size=10.0, lh=3.4, lh_s=3.08, measure_only=True,
)


def test_skill_gap_renders_between_categories_only():
    """skill_gap separates categories, so a one-category block has no gaps at
    all and the knob must not move the page height. It used to fire after the
    last category too, stacking a trailing gap onto the next section_gap."""
    r = ResumeRenderer()
    one = dict(SPARSE, skills={"Languages": ["Python", "SQL", "Go"]})

    _, tight_h = r._render_pdf(one, skill_gap=0.4, **_MEASURE_KNOBS)
    _, loose_h = r._render_pdf(one, skill_gap=3.0, **_MEASURE_KNOBS)
    assert tight_h == pytest.approx(loose_h, abs=1e-6)

    two = dict(SPARSE, skills={
        "Languages": ["Python", "SQL", "Go"],
        "Cloud & DevOps": ["AWS", "Docker"],
    })
    _, tight_2 = r._render_pdf(two, skill_gap=0.4, **_MEASURE_KNOBS)
    _, loose_2 = r._render_pdf(two, skill_gap=3.0, **_MEASURE_KNOBS)
    assert loose_2 - tight_2 == pytest.approx(2.6, abs=1e-6)  # exactly one gap


def test_skills_leading_capped_against_the_stretch():
    """Skills rows are dense comma lists — they must not take the body's full
    Pass-3 leading, or a sparse page renders them at ~5mm per line.

    Measured by differencing two resumes that vary only in category count, so
    the delta is purely skills lines.
    """
    r = ResumeRenderer()
    knobs = dict(_MEASURE_KNOBS, skill_gap=0.0, stretch=r._STRETCH_CAP)
    # Isolate the skills block so a stretched full resume cannot page-break
    # and hide the per-line delta (education keep-together starts page 2
    # at the top margin, which zeros overflow from extra skill rows).
    bare = {
        "contact": {"name": "Alex Rivera", "email": "a@example.com"},
        "summary": "", "experience": [], "education": [],
        "projects": [], "certifications": [],
    }
    one = dict(bare, skills={"A": ["Python"]})
    five = dict(bare, skills={k: ["Python"] for k in "ABCDE"})

    _, h1 = r._render_pdf(one, **knobs)
    _, h5 = r._render_pdf(five, **knobs)

    per_line = (h5 - h1) / 4
    # Uncapped this would be lh_s * stretch = 3.08 * 1.40 = 4.31mm.
    assert per_line == pytest.approx(r._MAX_SKILL_LEADING, abs=1e-6)


# ---------------------------------------------------------------------------
# Adaptive skills-section title
# ---------------------------------------------------------------------------

def test_skills_title_whitelist():
    from services.pdf_ats_check import extract_pdf_text

    r = ResumeRenderer()
    core = dict(SPARSE, skills_section_title="Core Skills")
    pdf_core, _ = r.generate_pdf(core)
    text, _ = extract_pdf_text(pdf_core)
    assert "CORE SKILLS" in text

    # Unknown/hostile values fall back to the default heading.
    junk = dict(SPARSE, skills_section_title="Best Stuff Ever")
    pdf_junk, _ = r.generate_pdf(junk)
    text, _ = extract_pdf_text(pdf_junk)
    assert "TECHNICAL SKILLS" in text
