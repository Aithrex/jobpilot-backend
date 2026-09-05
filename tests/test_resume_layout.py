"""Layout-core contract: no clipping, fill ceiling, Letter, parity, certs, CJK."""
from __future__ import annotations

import io
import json
import random
import re
from pathlib import Path

import pytest
from pypdf import PdfReader

from services.pdf_ats_check import check_pdf_parseability, extract_pdf_text
from services.resume_renderer import (
    DEFAULT_SPEC,
    PageResolution,
    ResumeRenderer,
    UnrenderableTextError,
)


def _bullet(i, n_chars=120):
    base = (
        f"Delivered measurable platform improvement number {i} by refactoring "
        f"the ingestion service and tuning PostgreSQL indexes for clinical workloads."
    )
    if n_chars <= len(base):
        return base[:n_chars]
    extra = " extra detail about scale, latency, reliability, and on-call response."
    text = base
    while len(text) < n_chars:
        text += extra
    return text[:n_chars]


def _resume(
    n_roles=2,
    bullets_per_role=4,
    n_projects=0,
    bullets_per_project=0,
    n_skills=4,
    skill_line_len=80,
    summary_len=200,
    n_edu=1,
    certs=None,
    company="Company",
    gpa="",
    name="Alex Rivera",
    bullet_chars=200,
):
    skills = {}
    for i in range(n_skills):
        pad = "x" * max(0, skill_line_len - 20)
        skills[f"Category {i}"] = [f"Skill{i}{pad}"[: max(4, skill_line_len // 2)]]
    summary = ("Backend engineer with production experience. " * 20)[:summary_len]
    return {
        "contact": {
            "name": name, "email": "a@example.com", "phone": "480-555-0100",
            "linkedin": "linkedin.com/in/alex", "github": "",
            "location": "Tempe, AZ", "portfolio": "",
        },
        "summary": summary,
        "skills": skills,
        "experience": [
            {
                "title": f"Engineer {r}",
                "company": f"{company} {r}" if n_roles > 1 else company,
                "location": "Tempe, AZ",
                "dates": f"January 202{r} - December 202{r}",
                "type": "Full-time",
                "bullets": [_bullet(f"{r}.{b}", bullet_chars) for b in range(bullets_per_role)],
            }
            for r in range(n_roles)
        ],
        "education": [
            {
                "degree": "B.S. Computer Science",
                "institution": f"State University {e}",
                "location": "Tempe, AZ",
                "dates": "May 2020",
                "gpa": gpa,
                "coursework": "Algorithms" if e == 0 else "",
            }
            for e in range(n_edu)
        ] if n_edu else [],
        "projects": [
            {
                "name": f"Project {p}", "dates": "",
                "bullets": [_bullet(f"p{p}.{b}", bullet_chars) for b in range(bullets_per_project)],
                "tech": "Python",
            }
            for p in range(n_projects)
        ],
        "certifications": certs or [],
    }


def _eight_by_six():
    # Matches the audited clip case: 8 roles × 6 long bullets overflow Letter
    # at the 9pt floor by more than MIN_SECOND_PAGE_FILL.
    return _resume(
        n_roles=8, bullets_per_role=6, n_projects=0,
        summary_len=400, bullet_chars=220,
    )


FIXTURES = {
    "sparse": _resume(2, 4),
    "new_grad": _resume(1, 3, n_projects=1, bullets_per_project=3),
    "senior": _resume(4, 5),
    "stress": _resume(3, 6, n_projects=3, bullets_per_project=4),
    "eight_by_six": _eight_by_six(),
}


def _usable(r: ResumeRenderer, pages: int, res: PageResolution | None = None) -> float:
    if res is not None:
        return pages * r.spec.usable_height(res.gaps.margin_top, res.gaps.margin_bottom)
    return pages * r.spec.usable_height()


@pytest.mark.parametrize("key", list(FIXTURES))
def test_no_clipping_ever(key):
    r = ResumeRenderer()
    resume = FIXTURES[key]
    pdf, res = r.generate_pdf(resume)
    assert isinstance(res, PageResolution)
    assert res.predicted_height_mm <= _usable(r, res.pages, res) + 0.75
    assert len(PdfReader(io.BytesIO(pdf)).pages) == res.pages
    assert res.pages >= 1
    assert res.outcome in ("fit_one", "fit_two", "trimmed_to_one", "trimmed_to_two",
                           "one_page_not_possible", "needs_shortening_choice")


def test_eight_by_six_does_not_clip_to_one_page():
    r = ResumeRenderer()
    pdf, res = r.generate_pdf(_eight_by_six())
    # Without silent truncation, dense content may need 2+ pages.
    assert res.pages >= 2
    assert res.outcome in ("fit_two", "trimmed_to_two", "needs_shortening_choice")
    assert res.predicted_height_mm <= _usable(r, res.pages, res) + 0.75
    text, pages = extract_pdf_text(pdf)
    assert pages == res.pages
    # Content is preserved (no silent drop of late roles).
    assert "Company 0" in text or "Engineer 0" in text
    assert "Company 7" in text or "Engineer 7" in text


@pytest.mark.parametrize("key", list(FIXTURES))
def test_round_trip_completeness(key):
    r = ResumeRenderer()
    resume = FIXTURES[key]
    pdf, res = r.generate_pdf(resume)
    text, _ = extract_pdf_text(pdf)
    collapsed = re.sub(r"\s+", " ", text)
    # Only assert bullets that the solver kept.
    kept = resume
    if res.dropped:
        # After trim, generate_pdf used a working copy; re-extract from PDF.
        pass
    for exp in kept.get("experience") or []:
        for b in exp.get("bullets") or []:
            needle = re.sub(r"\s+", " ", str(b)[:40]).strip()
            if not needle:
                continue
            # A trimmed bullet is listed in dropped previews.
            if any(needle[:30] in d.preview for d in res.dropped):
                continue
            assert needle in collapsed, f"missing bullet {needle!r} outcome={res.outcome}"


def test_page_two_fill_when_two_pages():
    r = ResumeRenderer()
    _pdf, res = r.generate_pdf(_eight_by_six())
    if res.pages == 2 and res.outcome == "fit_two":
        assert res.fill_ratio >= DEFAULT_SPEC.min_second_page_fill - 0.02


@pytest.mark.parametrize("key", list(FIXTURES))
def test_fill_ceiling(key):
    r = ResumeRenderer()
    _w, res = r._compute_pdf_layout(FIXTURES[key])
    if res.pages == 1:
        assert res.fill_ratio <= DEFAULT_SPEC.max_fill + 1e-6


def test_extraction_spacing_long_company_and_gpa():
    r = ResumeRenderer()
    resume = _resume(n_roles=1, bullets_per_role=3, gpa="3.80")
    resume["experience"][0]["company"] = (
        "International Business Machines Global Services"
    )
    resume["experience"][0]["location"] = "Tempe, AZ"
    resume["experience"][0]["dates"] = "January 2020 - December 2021"
    resume["education"][0]["gpa"] = "3.80"
    resume["education"][0]["dates"] = "May 2020"
    pdf, _res = r.generate_pdf(resume)
    text, _ = extract_pdf_text(pdf)
    collapsed = re.sub(r"\s+", " ", text)
    assert re.search(r"AZ\s+January 2020", collapsed), collapsed
    assert re.search(r"3\.80\s+May 2020", collapsed) or "3.80" in collapsed and "May 2020" in collapsed


def test_pdf_docx_share_layout_spec():
    spec = DEFAULT_SPEC
    assert spec.heading("summary") == "SUMMARY"
    assert spec.heading("certifications") == "CERTIFICATIONS"
    assert spec.section_order[-1] == "certifications"
    r = ResumeRenderer()
    assert r.spec.headings == spec.headings
    assert spec.pdf_font_family == "CharisSIL"
    assert spec.docx_font_family == "Georgia"
    assert spec.pdf_font_family != spec.docx_font_family
    from docx import Document
    docx_bytes = r.generate_docx(_resume(2, 4))
    doc = Document(io.BytesIO(docx_bytes))
    assert doc.styles["Normal"].font.name == spec.docx_font_family
    pdf, _ = r.generate_pdf(_resume(2, 4))
    assert spec.pdf_font_family.encode()[:6] in pdf or b"Charis" in pdf
    fe = Path(__file__).resolve().parents[2] / "jobpilot-frontend" / "src" / "lib" / "resumeLayoutSpec.ts"
    if fe.is_file():
        src = fe.read_text(encoding="utf-8")
        public = spec.to_public_dict()
        for key, value in public["headings"].items():
            assert value in src
        assert spec.pdf_font_family in src
        assert spec.docx_font_family in src
        assert spec.text_align in src


def test_certifications_present_and_absent():
    r = ResumeRenderer()
    with_certs = _resume(certs=["AWS Certified Cloud Practitioner", "CKA"])
    pdf, _ = r.generate_pdf(with_certs)
    text, _ = extract_pdf_text(pdf)
    assert "CERTIFICATIONS" in text
    assert "AWS Certified Cloud Practitioner" in text

    from docx import Document
    docx_bytes = r.generate_docx(with_certs)
    doc_text = "\n".join(p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs)
    assert "CERTIFICATIONS" in doc_text
    assert "AWS Certified Cloud Practitioner" in doc_text

    bare = _resume(certs=[])
    pdf2, _ = r.generate_pdf(bare)
    text2, _ = extract_pdf_text(pdf2)
    assert "CERTIFICATIONS" not in text2
    docx2 = r.generate_docx(bare)
    doc_text2 = "\n".join(p.text for p in Document(io.BytesIO(docx2)).paragraphs)
    assert "CERTIFICATIONS" not in doc_text2


def test_type_scale_is_absolute():
    spec = DEFAULT_SPEC
    assert spec.name_pt == 19.0
    assert spec.section_pt == 11.5
    assert spec.entry_pt == 10.2
    assert spec.contact_pt == 9.4
    r = ResumeRenderer()
    for key in FIXTURES:
        _w, res = r._compute_pdf_layout(FIXTURES[key])
        assert spec.section_pt - res.tier.body_pt >= spec.min_heading_body_gap_pt - 1e-9
        assert res.tier.body_pt <= 10.0
    # Headings do not shrink with body density.
    assert spec.name_pt == 19.0
    assert spec.section_pt == 11.5
    assert spec.entry_pt == 10.2


def test_cjk_name_has_no_question_marks():
    r = ResumeRenderer()
    resume = _resume(n_roles=1, bullets_per_role=2, name="李明")
    try:
        pdf, _res = r.generate_pdf(resume)
    except UnrenderableTextError:
        pytest.skip("NotoSansSC not installed in services/fonts/")
    text, _ = extract_pdf_text(pdf)
    assert "?" not in text
    assert "李明" in text or "李" in text


def test_no_render_warnings_on_happy_path(caplog):
    import logging
    caplog.set_level(logging.WARNING, logger="services.resume_renderer")
    r = ResumeRenderer()
    r.generate_pdf(_resume(2, 4))
    warnings = [
        rec.message for rec in caplog.records
        if rec.name.startswith("services.resume_renderer")
        and rec.levelno >= logging.WARNING
    ]
    assert warnings == [], warnings


def test_ats_single_page_bonus_only_for_fit_one():
    r = ResumeRenderer()
    resume = _resume(2, 4)
    pdf, res = r.generate_pdf(resume)
    assert res.outcome == "fit_one"
    honest = check_pdf_parseability(pdf, resume, {}, outcome="fit_one")
    trimmed = check_pdf_parseability(pdf, resume, {}, outcome="trimmed_to_one")
    assert honest["checks"]["single_page"] is True
    assert trimmed["checks"]["single_page"] is False
    assert honest["pdf_parseable_score"] == trimmed["pdf_parseable_score"] + 3


def test_resolution_json_shape():
    r = ResumeRenderer()
    _pdf, res = r.generate_pdf(_resume(2, 4))
    payload = res.to_dict()
    assert set(payload) >= {"pages", "outcome", "fill_ratio", "dropped"}
    json.dumps(payload)  # must be serializable


@pytest.mark.parametrize("seed", range(12))
def test_property_corpus_invariants(seed):
    rng = random.Random(seed)
    resume = _resume(
        n_roles=rng.randint(0, 5),
        bullets_per_role=rng.randint(0, 11),
        n_projects=rng.randint(0, 5),
        bullets_per_project=rng.randint(1, 5) if rng.random() > 0.3 else 0,
        n_skills=rng.randint(0, 10),
        skill_line_len=rng.randint(34, 200),
        summary_len=rng.randint(0, 874),
        n_edu=rng.randint(0, 4),
    )
    # Observed corpus includes 477-char bullets.
    if resume["experience"]:
        resume["experience"][0]["bullets"] = [
            _bullet(0, rng.choice([69, 195, 477]))
            for _ in range(max(1, len(resume["experience"][0]["bullets"]) or 1))
        ]
    r = ResumeRenderer()
    pdf, res = r.generate_pdf(resume)
    assert res.predicted_height_mm <= _usable(r, res.pages, res) + 0.75
    if res.pages == 1:
        assert res.fill_ratio <= DEFAULT_SPEC.max_fill + 1e-6
    text, pages = extract_pdf_text(pdf)
    assert pages == res.pages
    collapsed = re.sub(r"\s+", " ", text)
    for exp in resume.get("experience") or []:
        for b in exp.get("bullets") or []:
            needle = re.sub(r"\s+", " ", str(b)[:40]).strip()
            if not needle:
                continue
            if any(needle[:30] in d.preview for d in res.dropped):
                continue
            if needle in collapsed:
                continue
            # Empty-role / empty-summary fixtures can omit EXPERIENCE entirely.
            if not (resume.get("experience") or []):
                continue
            assert needle in collapsed


# ---------------------------------------------------------------------------
# Visual-style contract (Charis, margins, gap ordering)
# ---------------------------------------------------------------------------

# Snapshot of chosen body_pt under Charis SIL. A future font change should
# fail this explicitly rather than silently reflow.
DENSITY_TIER_SNAPSHOT = {
    "sparse": 10.0,
    "new_grad": 10.0,
    "senior": 9.4,
}


def test_density_ladder_charis_snapshot():
    r = ResumeRenderer()
    for key, expected in DENSITY_TIER_SNAPSHOT.items():
        _w, res = r._compute_pdf_layout(FIXTURES[key])
        assert res.tier.body_pt == expected, (
            f"{key}: expected {expected}pt, got {res.tier.body_pt}pt "
            f"(outcome={res.outcome} fill={res.fill_ratio:.3f})"
        )
    _w, stress = r._compute_pdf_layout(FIXTURES["stress"])
    assert stress.tier.body_pt <= 10.0
    _w, eight = r._compute_pdf_layout(FIXTURES["eight_by_six"])
    assert eight.pages >= 1


def test_docx_height_estimator_within_5_percent():
    r = ResumeRenderer()
    for key in ("sparse", "new_grad", "senior"):
        working, res = r._compute_pdf_layout(FIXTURES[key])
        est_pt = r._estimate_docx_height_pt(
            working,
            body_size=res.tier.body_pt,
            line_spacing=(DEFAULT_SPEC.body_lh_mm_at_ref * 72 / 25.4)
            / DEFAULT_SPEC.body_pt * res.stretch,
            section_gap=res.gaps.section_gap,
            entry_gap=res.gaps.entry_gap,
            skill_gap=res.gaps.skill_gap,
            header_gap=res.gaps.header_gap,
        )
        real_pt = res.predicted_height_mm * 72.0 / 25.4
        assert real_pt > 0
        err = abs(est_pt - real_pt) / real_pt
        assert err <= 0.05, (
            f"{key}: estimator off by {err:.1%} (est={est_pt:.1f} real={real_pt:.1f})"
        )


def test_side_margin_never_below_half_inch():
    r = ResumeRenderer()
    floor = r.spec.side_margin_dense_in * 25.4 - 1e-6
    for key in FIXTURES:
        _w, res = r._compute_pdf_layout(FIXTURES[key])
        assert res.gaps.margin_lr >= floor, key


def test_section_gap_exceeds_post_header_at_every_density():
    r = ResumeRenderer()
    for key in FIXTURES:
        _w, res = r._compute_pdf_layout(FIXTURES[key])
        assert res.gaps.section_gap > res.gaps.post_header, (
            f"{key}: section_gap {res.gaps.section_gap} <= post_header {res.gaps.post_header}"
        )
    assert r.spec.min_section_gap > r.spec.max_post_header


def test_heading_chrome_stays_in_range_across_densities():
    """Leftover whitespace may open between sections. post_header may compress
    toward min when needed to keep one page, but must stay within [min, max]."""
    r = ResumeRenderer()
    posts = []
    for key in FIXTURES:
        _w, res = r._compute_pdf_layout(FIXTURES[key])
        posts.append(res.gaps.post_header)
        assert res.gaps.post_header >= r.spec.min_post_header - 1e-9
        assert res.gaps.post_header <= r.spec.max_post_header + 1e-9
    assert r.spec.min_post_header < r.spec.max_post_header
    assert r.spec.contact_section_gap_mm == r.spec.min_section_gap
    # Rule sits in the middle of the reference heading chrome, not glued to text.
    assert abs(r.spec.heading_rule_gap_mm - r.spec.max_post_header) <= 0.40


def test_charis_embedded_and_tinos_fallback(monkeypatch):
    r = ResumeRenderer()
    pdf, _ = r.generate_pdf(_resume(1, 2))
    assert b"Charis" in pdf

    from fpdf import FPDF
    orig = FPDF.add_font

    def _hide_charis(self, family="", style="", fname="", uni=None, *a, **kw):
        if "Charis" in str(family) or "Charis" in str(fname):
            raise OSError("charis hidden for fallback test")
        return orig(self, family, style, fname, *a, **kw)

    monkeypatch.setattr(FPDF, "add_font", _hide_charis)
    pdf2, res = ResumeRenderer().generate_pdf(_resume(1, 2))
    assert pdf2.startswith(b"%PDF")
    assert res.pages >= 1
    assert b"Tinos" in pdf2 or b"Times" in pdf2


def test_text_align_is_left_and_never_justified():
    assert DEFAULT_SPEC.text_align == "L"
    renderer = (
        Path(__file__).resolve().parents[1]
        / "shared" / "python" / "services" / "resume_renderer.py"
    )
    src = renderer.read_text(encoding="utf-8")
    assert 'align="J"' not in src
    assert "align='J'" not in src


def test_bullet_glyphs_extract():
    r = ResumeRenderer()
    resume = _resume(n_roles=2, bullets_per_role=3, n_projects=1, bullets_per_project=2)
    expected = 2 * 3 + 2
    pdf, _res = r.generate_pdf(resume)
    text, _ = extract_pdf_text(pdf)
    count = text.count("•") + text.count("\u00b7")
    if count < expected:
        count = len(re.findall(r"(?:^|\n)\s*[-•\u00b7]", text))
    assert count >= expected, f"extracted {count} markers, expected {expected}\n{text[:800]}"


def test_headings_extract_as_single_tokens():
    r = ResumeRenderer()
    pdf, _ = r.generate_pdf(_resume(2, 3, n_projects=1, bullets_per_project=2))
    text, _ = extract_pdf_text(pdf)
    for heading in ("SUMMARY", "EXPERIENCE", "PROJECTS", "TECHNICAL SKILLS", "EDUCATION"):
        assert heading in text, heading
        lead = heading[0]
        rest = heading[1:].lstrip()
        assert not re.search(rf"{lead}\s+{re.escape(rest)}", text), heading


def test_sixty_char_company_keeps_location_date_space():
    r = ResumeRenderer()
    resume = _resume(n_roles=1, bullets_per_role=3)
    resume["experience"][0]["company"] = "A" * 60
    resume["experience"][0]["location"] = "Tempe, AZ"
    resume["experience"][0]["dates"] = "January 2020 - December 2021"
    pdf, _ = r.generate_pdf(resume)
    text, _ = extract_pdf_text(pdf)
    collapsed = re.sub(r"\s+", " ", text)
    assert re.search(r"AZ\s+January 2020", collapsed), collapsed


def test_skills_markdown_does_not_print_unknown():
    r = ResumeRenderer()
    resume = _resume(n_roles=1, bullets_per_role=2, n_projects=1, bullets_per_project=2)
    resume["skills"] = {
        "": ["Backend"],
        "Concepts & Methodologies": ["Microservices", "Full-Stack", "Java 17"],
        "Languages": ["Python", "SQL"],
    }
    resume["projects"][0]["name"] = "Document Intelligence RAG Platform"
    resume["projects"][0]["tech"] = "ZZZUNIQUE_TECH_TOKEN, pgvector"
    resume["education"][0] = {
        "degree": "Master of Science, Information Technology",
        "institution": "Arizona State University",
        "location": "Tempe, AZ",
        "dates": "January 2024 - December 2025",
        "gpa": "",
        "coursework": "",
    }
    pdf, _ = r.generate_pdf(resume)
    text, _ = extract_pdf_text(pdf)
    collapsed = re.sub(r"\s+", " ", text)
    assert "UNKNOWN" not in text
    assert "< UNKNOWN >" not in collapsed
    assert "ZZZUNIQUE_TECH_TOKEN" not in text
    assert "Java" in collapsed
    assert "Python" in collapsed
    assert "Microservices" in collapsed
    assert "Arizona State University" in collapsed
    assert "Master of Science" in collapsed
    # Role labels must not appear as skill items.
    assert not re.search(r":\s*Backend\b", collapsed)
    assert "Full-Stack" not in collapsed


def test_visual_snapshots():
    """Render three fixtures to PNG. Compared against committed goldens when
    pypdfium2 is installed; otherwise skipped so CI still runs."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        pytest.skip("pypdfium2 not installed — PNG snapshots are local/dev")

    snap_dir = Path(__file__).resolve().parent / "snapshots" / "resume"
    snap_dir.mkdir(parents=True, exist_ok=True)
    cases = {
        "sparse_new_grad": FIXTURES["new_grad"],
        "median_tailored": FIXTURES["senior"],
        "dense_5_role": _resume(n_roles=5, bullets_per_role=5, bullet_chars=180),
    }
    r = ResumeRenderer()
    for name, resume in cases.items():
        pdf, _res = r.generate_pdf(resume)
        doc = pdfium.PdfDocument(pdf)
        page = doc[0]
        pil = page.render(scale=2).to_pil()
        golden = snap_dir / f"{name}.png"
        if not golden.exists():
            pil.save(golden)
            pytest.fail(f"wrote new golden {golden}; re-run to confirm")
        from PIL import Image, ImageChops
        existing = Image.open(golden)
        if pil.size != existing.size:
            pil.save(snap_dir / f"{name}.new.png")
            pytest.fail(f"{name}: size {pil.size} != golden {existing.size}")
        diff = ImageChops.difference(pil.convert("RGB"), existing.convert("RGB"))
        extrema = diff.getextrema()
        max_delta = max(ch[1] for ch in extrema)
        if max_delta > 12:
            pil.save(snap_dir / f"{name}.new.png")
            pytest.fail(f"{name}: visual delta {max_delta} exceeds tolerance")
