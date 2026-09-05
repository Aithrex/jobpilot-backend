"""Package 3: ordered artifact extraction and word-boundary matching."""
from __future__ import annotations

import io

import pytest
from docx import Document

from services.artifact_assessment import (
    assess_artifact,
    extract_docx_ordered,
    extract_docx_ordered_legacy_concat_bug,
    format_cache_key,
    normalize_preserve_boundaries,
    token_boundary_match,
)
from tests.fixtures.assessment.cases import SQL_NOSQL_JD, SQL_NOSQL_RESUME


def test_sql_does_not_match_nosql():
    hay = "Operated MongoDB NoSQL clusters for product analytics."
    assert token_boundary_match(hay, "NoSQL") is True
    assert token_boundary_match(hay, "SQL") is False
    assert token_boundary_match("Wrote SQL against the warehouse", "SQL") is True


def test_normalize_keeps_boundaries():
    assert "sql" in normalize_preserve_boundaries("NoSQL")
    # Must not collapse to substring-only haystack without separators
    assert normalize_preserve_boundaries("No SQL") == "no sql"


def _docx_table_then_para_order() -> bytes:
    """Build DOCX where a table appears BEFORE a later paragraph in body order."""
    doc = Document()
    doc.add_paragraph("INTRO_PARAGRAPH")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TABLE_CELL_ALPHA"
    doc.add_paragraph("AFTER_TABLE_PARAGRAPH")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_preserves_document_order():
    data = _docx_table_then_para_order()
    parsed = extract_docx_ordered(data)
    texts = [b.text for b in parsed.blocks]
    assert "INTRO_PARAGRAPH" in texts
    assert "TABLE_CELL_ALPHA" in texts
    assert "AFTER_TABLE_PARAGRAPH" in texts
    # Table cell must appear between the two paragraphs in body order.
    i_intro = texts.index("INTRO_PARAGRAPH")
    i_table = texts.index("TABLE_CELL_ALPHA")
    i_after = texts.index("AFTER_TABLE_PARAGRAPH")
    assert i_intro < i_table < i_after


def test_legacy_concat_puts_table_after_all_paragraphs():
    data = _docx_table_then_para_order()
    legacy = extract_docx_ordered_legacy_concat_bug(data)
    # Bug: all paragraphs first, then tables — AFTER comes before TABLE.
    assert legacy.index("AFTER_TABLE_PARAGRAPH") < legacy.index("TABLE_CELL_ALPHA")


def test_pdf_docx_cache_keys_differ():
    k_pdf = format_cache_key(fmt="pdf", content_hash="c", data_hash="d")
    k_docx = format_cache_key(fmt="docx", content_hash="c", data_hash="d")
    assert k_pdf != k_docx


def test_assess_artifact_docx_sql_nosql_fixture():
    """Skills list NoSQL must not satisfy a SQL-only intended target check wrongly.

    Here we only assert boundary matching on extracted text; requirement
    evaluation is Package 4.
    """
    doc = Document()
    doc.add_paragraph("Data engineer working with NoSQL document stores.")
    doc.add_paragraph("Skills: NoSQL, MongoDB")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    result = assess_artifact(data, SQL_NOSQL_RESUME, fmt="docx")
    assert result["status"] in {"passed", "issues", "failed"}
    assert result["format"] == "docx"
    assert token_boundary_match(result["raw_text"], "SQL") is False
    assert token_boundary_match(result["raw_text"], "NoSQL") is True


def test_assess_artifact_error_and_empty():
    assert assess_artifact(None, {}, error="boom")["status"] == "failed"
    assert assess_artifact(None, {})["status"] == "unchecked"
