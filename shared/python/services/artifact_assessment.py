"""Ordered PDF/DOCX artifact assessment (assessment-v2 Package 3).

Preserves document order and word boundaries. ``SQL`` must not match
``NoSQL``. DOCX must traverse body elements in order (not all paragraphs
then all tables). Format-specific identity — never reuse PDF checks as DOCX.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from typing import Any, Dict, List, Literal, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)

EXTRACTOR_VERSION = "artifact-v1"
BlockKind = Literal[
    "paragraph", "table_cell", "heading", "list_item", "page_text", "unknown",
]


@dataclass
class DocumentBlock:
    """One ordered unit of extracted document text."""

    index: int
    kind: BlockKind
    text: str
    page: Optional[int] = None
    section_hint: str = ""
    entity_hint: str = ""
    offset_start: int = 0
    offset_end: int = 0

    def as_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class ParsedArtifact:
    format: str
    bytes_hash: str
    extractor_version: str = EXTRACTOR_VERSION
    renderer_version: str = ""
    page_count: Optional[int] = None
    blocks: List[DocumentBlock] = field(default_factory=list)
    raw_text: str = ""
    normalized_text: str = ""
    extraction_complete: bool = True
    issues: List[str] = field(default_factory=list)

    def as_dict(self) -> Dict[str, Any]:
        return {
            "format": self.format,
            "bytes_hash": self.bytes_hash,
            "extractor_version": self.extractor_version,
            "renderer_version": self.renderer_version,
            "page_count": self.page_count,
            "blocks": [b.as_dict() for b in self.blocks],
            "raw_text": self.raw_text,
            "normalized_text": self.normalized_text,
            "extraction_complete": self.extraction_complete,
            "issues": list(self.issues),
        }


def normalize_preserve_boundaries(text: str) -> str:
    """Normalize Unicode/spacing but keep word boundaries (unlike legacy strip-all)."""
    if not text:
        return ""
    t = unicodedata.normalize("NFKC", text)
    t = t.replace("\u2013", "-").replace("\u2014", "-").replace("\u00a0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.casefold().strip()


def token_boundary_match(haystack: str, needle: str) -> bool:
    """True when needle appears as a whole token/phrase, not a substring of another word.

    ``SQL`` does not match inside ``NoSQL``. Phrases match with flexible internal
    whitespace. Empty needle is never a match.
    """
    n = normalize_preserve_boundaries(needle)
    h = normalize_preserve_boundaries(haystack)
    if not n or not h:
        return False
    # Escape and allow flexible whitespace between needle tokens.
    parts = [re.escape(p) for p in n.split() if p]
    if not parts:
        return False
    inner = r"\s+".join(parts)
    pattern = rf"(?<![a-z0-9_]){inner}(?![a-z0-9_])"
    return re.search(pattern, h) is not None


def bytes_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def format_cache_key(*, fmt: str, content_hash: str, data_hash: str,
                     renderer_version: str = "") -> str:
    """Format-specific cache identity — PDF and DOCX never share a key."""
    return hashlib.sha256(
        f"{fmt}|{content_hash}|{data_hash}|{renderer_version}|{EXTRACTOR_VERSION}".encode()
    ).hexdigest()


def _assign_offsets(blocks: List[DocumentBlock]) -> str:
    parts: List[str] = []
    cursor = 0
    for b in blocks:
        b.offset_start = cursor
        parts.append(b.text)
        cursor += len(b.text)
        b.offset_end = cursor
        parts.append("\n")
        cursor += 1
    return "".join(parts).rstrip("\n")


def extract_docx_ordered(data: bytes) -> ParsedArtifact:
    """Traverse DOCX body in document order (paragraphs and tables interleaved)."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(data))
    blocks: List[DocumentBlock] = []
    body = doc.element.body
    idx = 0
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if text.strip():
                blocks.append(DocumentBlock(index=idx, kind="paragraph", text=text.strip()))
                idx += 1
        elif tag == qn("w:tbl"):
            for row in child.iterchildren():
                if row.tag != qn("w:tr"):
                    continue
                for cell in row.iterchildren():
                    if cell.tag != qn("w:tc"):
                        continue
                    cell_text = "".join(
                        node.text or "" for node in cell.iter(qn("w:t"))
                    ).strip()
                    if cell_text:
                        blocks.append(DocumentBlock(
                            index=idx, kind="table_cell", text=cell_text,
                        ))
                        idx += 1
    raw = _assign_offsets(blocks)
    return ParsedArtifact(
        format="docx",
        bytes_hash=bytes_hash(data),
        blocks=blocks,
        raw_text=raw,
        normalized_text=normalize_preserve_boundaries(raw),
        extraction_complete=True,
    )


def extract_docx_ordered_legacy_concat_bug(data: bytes) -> str:
    """Historical wrong order: all paragraphs then all tables (for regression tests)."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )


def extract_pdf_ordered(data: bytes) -> ParsedArtifact:
    """One block per PDF page in page order."""
    from services.pdf_ats_check import extract_pdf_text
    # pypdf does not expose rich structure here; page-level order is preserved.
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        blocks: List[DocumentBlock] = []
        for i, page in enumerate(reader.pages):
            try:
                text = (page.extract_text() or "").strip()
            except Exception:
                text = ""
            if text:
                blocks.append(DocumentBlock(
                    index=i, kind="page_text", text=text, page=i + 1,
                ))
        raw = _assign_offsets(blocks)
        return ParsedArtifact(
            format="pdf",
            bytes_hash=bytes_hash(data),
            page_count=len(reader.pages),
            blocks=blocks,
            raw_text=raw,
            normalized_text=normalize_preserve_boundaries(raw),
            extraction_complete=bool(raw.strip()),
            issues=[] if raw.strip() else ["No extractable PDF text"],
        )
    except Exception as exc:
        text, pages = extract_pdf_text(data)
        blocks = []
        if text.strip():
            blocks = [DocumentBlock(index=0, kind="page_text", text=text.strip(), page=1)]
        raw = text
        return ParsedArtifact(
            format="pdf",
            bytes_hash=bytes_hash(data),
            page_count=pages or None,
            blocks=blocks,
            raw_text=raw,
            normalized_text=normalize_preserve_boundaries(raw),
            extraction_complete=bool(raw.strip()),
            issues=[f"PDF extraction degraded: {exc}"],
        )


def extract_artifact(data: bytes, *, fmt: str) -> ParsedArtifact:
    fmt_n = (fmt or "pdf").strip().lower()
    if fmt_n == "docx":
        return extract_docx_ordered(data)
    if fmt_n == "pdf":
        return extract_pdf_ordered(data)
    raise ValueError(f"unsupported format: {fmt}")


def intended_targets(resume: Dict[str, Any]) -> List[Tuple[str, str]]:
    """Labels + expected strings from structured resume content."""
    targets: List[Tuple[str, str]] = []
    for key, val in (resume.get("contact") or {}).items():
        if isinstance(val, str) and val.strip() and key in {"name", "email", "phone"}:
            targets.append((f"contact.{key}", val))
    if resume.get("summary"):
        targets.append(("summary", str(resume["summary"])))
    for section in ("experience", "education", "projects"):
        for i, row in enumerate(resume.get(section) or []):
            for field in ("company", "title", "dates", "institution", "degree", "name"):
                if row.get(field):
                    targets.append((f"{section}.{i}.{field}", str(row[field])))
            for j, b in enumerate(row.get("bullets") or []):
                targets.append((f"{section}.{i}.bullets.{j}", str(b)))
    try:
        from schemas.resume_schemas import flatten_skills
        targets.extend((f"skills.{i}", s) for i, s in enumerate(flatten_skills(resume)))
    except Exception:
        pass
    for i, c in enumerate(resume.get("certifications") or []):
        targets.append((f"certifications.{i}", str(c)))
    return targets


def assess_artifact(
    data: bytes | None,
    resume: Dict[str, Any],
    *,
    fmt: str = "pdf",
    error: str | None = None,
    renderer_version: str = "",
) -> Dict[str, Any]:
    """Return artifact stage payload: status, issues, ordered blocks, checks."""
    from services.tailoring_policy import content_hash

    out: Dict[str, Any] = {
        "format": fmt,
        "status": "unchecked",
        "issues": [],
        "checks": {},
        "bytes_hash": None,
        "content_hash": content_hash(resume) if resume else "",
        "extractor_version": EXTRACTOR_VERSION,
        "renderer_version": renderer_version,
        "cache_key": None,
        "blocks": [],
        "raw_text": "",
        "page_count": None,
        "extraction_complete": False,
    }
    if error:
        out.update(status="failed", issues=["Document could not be rendered."])
        return out
    if not data:
        return out

    try:
        parsed = extract_artifact(data, fmt=fmt)
    except Exception as exc:
        logger.warning("artifact extract failed: %s", exc)
        out.update(status="failed", issues=["Document text could not be extracted."])
        return out

    out["bytes_hash"] = parsed.bytes_hash
    out["blocks"] = [b.as_dict() for b in parsed.blocks]
    out["raw_text"] = parsed.raw_text
    out["page_count"] = parsed.page_count
    out["extraction_complete"] = parsed.extraction_complete
    out["cache_key"] = format_cache_key(
        fmt=fmt,
        content_hash=out["content_hash"],
        data_hash=parsed.bytes_hash,
        renderer_version=renderer_version,
    )

    checks: Dict[str, bool] = {
        "text_extracted": bool(parsed.raw_text.strip()),
        "no_replacement_glyphs": "\ufffd" not in parsed.raw_text,
        "ordered_blocks": len(parsed.blocks) > 0,
    }
    for label, value in intended_targets(resume):
        checks[label] = token_boundary_match(parsed.raw_text, value)

    out["checks"] = checks
    out["issues"] = [f"Export check failed: {k}" for k, ok in checks.items() if not ok]
    out["issues"].extend(parsed.issues)
    if not checks.get("text_extracted"):
        out["status"] = "failed"
    elif out["issues"]:
        out["status"] = "issues"
    else:
        out["status"] = "passed"
    return out


def text_contains_capability(text: str, capability: str) -> bool:
    """Public helper for Package 4 observers — word-boundary safe."""
    return token_boundary_match(text, capability)
