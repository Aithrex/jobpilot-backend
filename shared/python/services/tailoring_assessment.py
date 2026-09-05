"""Read-only, version-bound document assessment. No vendor simulations."""
from __future__ import annotations

import copy
import hashlib
import json
import re

from services.tailoring_policy import VERSION, REVIEW_RULES, content, content_hash, digest
from services.tailoring_planner import contains, semantic_support
from utils.llm_safety import guarded_system, wrap_untrusted, redact_pii_for_model


def _normalized(text: str) -> str:
    return re.sub(r"\s+", "", text).casefold().replace("\u2013", "-").replace("\u2014", "-")


def check_export(data: bytes | None, resume: dict, *, error: str | None = None,
                 format: str = "pdf") -> tuple[dict, str]:
    artifact = {"format": format, "content_hash": content_hash(resume), "status": "unchecked",
                "bytes_hash": None, "checks": {}, "issues": []}
    if error:
        artifact.update(status="failed", issues=["Document could not be rendered."])
        return artifact, ""
    if not data:
        return artifact, ""
    artifact["bytes_hash"] = hashlib.sha256(data).hexdigest()
    try:
        if format == "docx":
            import io
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] +
                             [c.text for t in doc.tables for r in t.rows for c in r.cells])
        else:
            from services.pdf_ats_check import extract_pdf_text
            text, _ = extract_pdf_text(data)
        norm = _normalized(text)
        checks = {"text_extracted": bool(text.strip()), "no_replacement_glyphs": "\ufffd" not in text}
        targets = []
        for key, val in (resume.get("contact") or {}).items():
            if isinstance(val, str) and val.strip() and key in {"name", "email", "phone"}:
                targets.append(("contact." + key, val))
        for section in ("experience", "education", "projects"):
            for i, row in enumerate(resume.get(section) or []):
                for field in ("company", "title", "dates", "institution", "degree", "name"):
                    if row.get(field):
                        targets.append((f"{section}.{i}.{field}", row[field]))
                for j, b in enumerate(row.get("bullets") or []):
                    targets.append((f"{section}.{i}.bullets.{j}", b))
        if resume.get("summary"):
            targets.append(("summary", resume["summary"]))
        from schemas.resume_schemas import flatten_skills
        targets.extend((f"skills.{i}", s) for i, s in enumerate(flatten_skills(resume)))
        targets.extend((f"certifications.{i}", c) for i, c in enumerate(resume.get("certifications") or []))
        for label, value in targets:
            checks[label] = _normalized(str(value)) in norm
        artifact["checks"] = checks
        artifact["issues"] = [f"Export check failed: {k}" for k, ok in checks.items() if not ok]
        artifact["status"] = "passed" if all(checks.values()) else "issues"
        return artifact, text
    except Exception:
        artifact.update(status="failed", issues=["Document text could not be extracted."])
        return artifact, ""


def reader_review(visible_text: str, jd: dict) -> dict:
    from services.claude_client import claude_json, CLAUDE_PRO
    if not visible_text:
        return {"status": "unverified", "findings": [], "reason": "No extracted document text."}
    # Skip oversized review rather than pretend a prefix is a full-document read.
    if len(visible_text) > 45000:
        return {"status": "unverified", "findings": [], "reason": "Document exceeds review budget."}
    try:
        job = {k: jd.get(k) for k in ("job_title", "responsibilities", "required_skills", "preferred_skills", "qualifications")}
        response = claude_json(
            prompt=wrap_untrusted("visible_resume", redact_pii_for_model(visible_text)) + "\n" +
                   wrap_untrusted("job", json.dumps(job)),
            system=guarded_system(REVIEW_RULES), model=CLAUDE_PRO, temperature=0,
            max_tokens=2500, schema={"apparent_role": str, "apparent_level": str,
                                    "findings": [{"kind": str, "quote": str, "explanation": str}]})
        findings = [r for r in response.get("findings") or [] if isinstance(r, dict)
                    and r.get("quote") and _normalized(r["quote"]) in _normalized(visible_text)
                    and r.get("kind") in {"strength", "concern", "clarity", "over_tailoring", "hidden_fit"}]
        return {"status": "passed", "apparent_role": response.get("apparent_role", ""),
                "apparent_level": response.get("apparent_level", ""), "findings": findings}
    except Exception:
        return {"status": "unverified", "findings": [], "reason": "Reader review unavailable."}


def assess(resume: dict, jd: dict, data: bytes | None = None, error: str | None = None,
           *, run_reader: bool = True) -> dict:
    run = resume.get("tailoring_run") or {}
    current_hash = content_hash(resume)
    stale = run.get("content_hash") != current_hash or run.get("jd_hash") != digest(jd)
    artifact, text = check_export(data, resume, error=error)
    provenance = run.get("provenance") or []
    reqs = copy.deepcopy(run.get("requirements") or [])
    for req in reqs:
        req["locations"] = [p.get("entity") for p in provenance if p.get("fact_id") in req.get("fact_ids", [])]
        linked = [p.get("text", "") for p in provenance if p.get("fact_id") in req.get("fact_ids", [])]
        req["visible"] = any(semantic_support(req["text"], t) or any(contains(a, t) for a in req.get("alternatives") or []) for t in linked)
        req["visible"] |= any(contains(a, json.dumps(resume.get("skills", {}))) for a in req.get("alternatives") or [])
        req["visible"] |= any(contains(a, json.dumps([resume.get("education", []), resume.get("certifications", [])]))
                               for a in req.get("alternatives") or [])
        if stale:
            req["state"] = "needs_clarification"
            req["locations"] = []
    existing = resume.get("document_assessment") or {}
    cached_reader = existing.get("reader_review")
    reuse = (existing.get("content_hash") == current_hash and existing.get("jd_hash") == digest(jd)
             and (existing.get("artifact") or {}).get("bytes_hash") == artifact.get("bytes_hash"))
    reader = cached_reader if reuse and cached_reader else (reader_review(text, jd) if run_reader else {
        "status": "unverified", "findings": [], "reason": "Reader review not requested."})
    manual = any(p.get("state") == "user_edit_unverified" for p in provenance) or bool((resume.get("tailoring_preferences") or {}).get("summary_locked"))
    return {"version": VERSION, "content_hash": current_hash, "jd_hash": digest(jd),
            "artifact": artifact, "requirements": reqs, "reader_review": reader,
            "source_status": "stale" if stale else "user_edits_unverified" if manual else "source_linked",
            "questions": run.get("questions", []), "analysis_completeness": run.get("analysis_completeness", "unknown"),
            "ready": not stale and not manual and artifact["status"] == "passed" and reader["status"] == "passed"
                     and not any(f.get("kind") in {"concern", "over_tailoring"} for f in reader.get("findings", []))}


def finalize(resume: dict, jd: dict, renderer) -> dict:
    existing = resume.get("document_assessment") or {}
    if (existing.get("content_hash") == content_hash(resume)
            and existing.get("jd_hash") == digest(jd)):
        return resume
    data, error = None, None
    try:
        data, _ = renderer.generate_pdf(content(resume))
    except Exception:
        error = "render_failed"
    result = copy.deepcopy(resume)
    result["document_assessment"] = assess(result, jd, data, error)
    return result


def score_payload(resume: dict, jd: dict, data: bytes | None, error: str | None) -> dict:
    assessment = assess(resume, jd, data, error)
    # Numeric legacy aliases are intentionally not populated with estimates.
    # New clients render document_assessment; absence means unknown to adapters.
    return {"schema_version": 3, "assessment_version": VERSION,
            "document_assessment": assessment, "overall": None,
            "suggestions": [r.get("explanation", "") for r in assessment["reader_review"].get("findings", []) if r.get("kind") != "strength"],
            "strengths": [r.get("explanation", "") for r in assessment["reader_review"].get("findings", []) if r.get("kind") == "strength"],
            "missing_keywords": [r["text"] for r in assessment["requirements"] if r["state"] == "not_evidenced"]}
